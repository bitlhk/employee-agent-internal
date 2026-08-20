#!/usr/bin/env python3
"""Generate small, parseable Office files for the fictional GRACE demo corpus."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path


DISCLAIMER = "本文件为 GRACE 演示银行虚构资产，仅用于产品演示与测试。"
FIXED_TIME = datetime(2026, 8, 20, 0, 0, 0, tzinfo=timezone.utc)


def write_docx(target: Path, title: str, summary: str, asset_id: str, version: str, department: str) -> None:
    from docx import Document

    document = Document()
    document.core_properties.author = "GRACE Demo"
    document.core_properties.last_modified_by = "GRACE Demo"
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.add_heading(title, level=1)
    document.add_paragraph(DISCLAIMER)
    document.add_heading("适用范围", level=2)
    document.add_paragraph(f"{summary} 使用时必须核验当前有效版本、数据时点和岗位权限。")
    document.add_heading("工作要求", level=2)
    for text in [
        "先确认任务所需事实是否完整，不以长期记忆替代实时业务数据。",
        "只采用与当前岗位、机构和用途相符的企业依据。",
        "涉及写入或业务状态变更时，必须经过策略校验和必要的人工确认。",
    ]:
        document.add_paragraph(text, style="List Number")
    table = document.add_table(rows=4, cols=2)
    for row, values in zip(table.rows, [
        ("资产编号", asset_id),
        ("版本", version or "待确认"),
        ("责任部门", department),
        ("说明", summary),
    ]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(target)


def write_xlsx(target: Path, title: str, summary: str, asset_id: str, version: str, department: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.properties.creator = "GRACE Demo"
    workbook.properties.lastModifiedBy = "GRACE Demo"
    workbook.properties.created = FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = FIXED_TIME.replace(tzinfo=None)
    sheet = workbook.active
    sheet.title = "演示资产"
    rows = [
        ["项目", "演示值", "治理说明"],
        ["资产名称", title, "虚构演示资产"],
        ["资产编号", asset_id, "用于 Ground Truth 对照"],
        ["版本", version or "待确认", "版本字段需由企业确认"],
        ["责任部门", department, "Owner 属于企业事实"],
        ["检查项一", "数据时点", "业务数据应实时获取"],
        ["检查项二", "有效版本", "过期依据不得进入 Context"],
        ["检查项三", "人工确认", "副作用动作必须确认"],
        ["说明", summary, DISCLAIMER],
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C7000B")
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 48
    workbook.save(target)


def write_pptx(target: Path, title: str, summary: str, asset_id: str, version: str, department: str) -> None:
    from pptx import Presentation

    presentation = Presentation()
    presentation.core_properties.author = "GRACE Demo"
    presentation.core_properties.last_modified_by = "GRACE Demo"
    presentation.core_properties.created = FIXED_TIME
    presentation.core_properties.modified = FIXED_TIME
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = summary
    for text in [
        "先核验任务事实与当前有效依据",
        "副作用动作经过策略校验和人工确认",
        DISCLAIMER,
    ]:
        body.add_paragraph().text = text
    detail = presentation.slides.add_slide(presentation.slide_layouts[1])
    detail.shapes.title.text = "资产治理信息"
    detail_body = detail.placeholders[1].text_frame
    detail_body.text = f"资产编号：{asset_id}"
    detail_body.add_paragraph().text = f"版本：{version or '待确认'}"
    detail_body.add_paragraph().text = f"责任部门：{department}"
    presentation.save(target)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=["docx", "xlsx", "pptx"], required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--title", required=True)
    parser.add_argument("--summary", required=True)
    parser.add_argument("--asset-id", required=True)
    parser.add_argument("--version", default="")
    parser.add_argument("--department", required=True)
    args = parser.parse_args()

    target = Path(args.output)
    target.parent.mkdir(parents=True, exist_ok=True)
    writers = {"docx": write_docx, "xlsx": write_xlsx, "pptx": write_pptx}
    writers[args.format](target, args.title, args.summary, args.asset_id, args.version, args.department)


if __name__ == "__main__":
    main()
